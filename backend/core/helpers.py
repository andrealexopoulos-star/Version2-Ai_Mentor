"""
Helper functions — extracted from server.py.
Includes: HTML stripping, website fetching, search, file parsing, auth utilities.
"""
import os
import io
import re
import logging
from typing import Dict, Any, List, Optional
from datetime import datetime, timezone, timedelta

import httpx
import jwt
import bcrypt
from fastapi import Depends, HTTPException
from fastapi.security import HTTPBearer, HTTPAuthorizationCredentials

logger = logging.getLogger(__name__)

# Document parsing imports (optional)
try:
    import PyPDF2
    from docx import Document as DocxDocument
    import openpyxl
except ImportError:
    PyPDF2 = None
    DocxDocument = None
    openpyxl = None

SERPER_API_KEY = os.environ.get("SERPER_API_KEY")  # retained: still referenced by health/quota probes
PERPLEXITY_API_KEY = os.environ.get("PERPLEXITY_API_KEY")
PERPLEXITY_SEARCH_MODEL = os.environ.get("PERPLEXITY_SEARCH_MODEL", "sonar")
JWT_SECRET = os.environ.get("JWT_SECRET_KEY", "")
JWT_ALGORITHM = "HS256"
JWT_EXPIRATION_HOURS = 24

MAX_FILE_SIZE = 10 * 1024 * 1024  # 10MB


# ==================== HTML / WEB HELPERS ====================

def strip_html_to_text(html: str) -> str:
    from bs4 import BeautifulSoup
    soup = BeautifulSoup(html or "", "html.parser")
    for tag in soup(["script", "style", "noscript", "svg"]):
        tag.decompose()
    text = soup.get_text("\n")
    lines = [ln.strip() for ln in text.splitlines()]
    lines = [ln for ln in lines if ln]
    return "\n".join(lines)


async def fetch_website_text(url: str) -> str:
    if not url:
        return ""
    url = url.strip()
    if not url.startswith("http://") and not url.startswith("https://"):
        url = f"https://{url}"
    async with httpx.AsyncClient(timeout=15, follow_redirects=True) as client:
        resp = await client.get(url, headers={
            "User-Agent": "Mozilla/5.0 (compatible; StrategySquadBot/1.0)"
        })
        if resp.status_code >= 400:
            return ""
        return strip_html_to_text(resp.text)


def compute_missing_profile_fields(profile_patch: Dict[str, Any]) -> List[str]:
    essentials = ["business_name", "industry", "business_type", "target_country", "retention_known"]
    missing = []
    for f in essentials:
        v = profile_patch.get(f)
        if v is None or (isinstance(v, str) and not v.strip()):
            missing.append(f)
    return missing


# ==================== SEARCH ====================

async def serper_search(query: str, gl: str = "au", hl: str = "en", num: int = 5) -> Dict[str, Any]:
    """Return {results: [{title, link, snippet, position}], error: str|None}.

    Function name preserved for back-compat with all callers. Implementation
    delegates to Perplexity sonar (Serper retired 2026-05-05 13041978 — Serper
    credits exhausted, switching primary search provider to Perplexity which
    is already proven across 5 production edge functions).

    Per BIQc_PLATFORM_CONTRACT_SECURE_NO_SILENT_FAILURE_v2 + zero-401: any
    upstream non-200 surfaces as an explicit error in the return value AND
    in the per-scan enrichment_traces row. We never silently return empty
    results on provider failure.
    """
    # Resolve active scan context. None outside the scan path → no trace.
    try:
        from core.enrichment_trace import (
            get_active_scan_id, get_active_user_id, arecord_provider_trace,
        )
        _scan_id = get_active_scan_id()
        _user_id = get_active_user_id()
    except Exception:
        _scan_id = None
        _user_id = None

    request_summary = {
        "provider_url": "api.perplexity.ai/chat/completions",
        "query_chars": len(query or ""),
        "gl": gl, "hl": hl, "num": num,
        "model": PERPLEXITY_SEARCH_MODEL,
    }
    import time as _t
    _t0 = _t.perf_counter()

    if not PERPLEXITY_API_KEY:
        if _scan_id:
            await arecord_provider_trace(
                scan_id=_scan_id, provider="perplexity", user_id=_user_id,
                http_status=503,
                latency_ms=int((_t.perf_counter() - _t0) * 1000),
                request_summary=request_summary,
                response_summary={"ok": False, "code": "PROVIDER_KEY_MISSING"},
                error="PROVIDER_KEY_MISSING : PERPLEXITY_API_KEY not configured",
                sanitiser_applied=True,
            )
        return {"results": [], "error": "PERPLEXITY_API_KEY not configured"}

    region_hint = ""
    if (gl or "").lower() in ("au", "aus", "australia"):
        region_hint = " Bias to Australian (.au) results when relevant."

    user_prompt = (
        f"Return ONLY a valid JSON array (no prose, no markdown fences) of the top {num} "
        f"most-relevant web search results for the query below. Each item MUST be an object "
        f"with EXACTLY these keys: title (string), url (string), snippet (string up to 240 chars).{region_hint}\n\n"
        f"QUERY: {query}"
    )

    payload = {
        "model": PERPLEXITY_SEARCH_MODEL,
        "messages": [
            {"role": "system", "content": "You are a web search API. Return ONLY a JSON array as specified. No prose, no preamble, no code fences."},
            {"role": "user", "content": user_prompt},
        ],
        "temperature": 0.0,
        "max_tokens": 1200,
    }
    headers = {
        "Authorization": f"Bearer {PERPLEXITY_API_KEY}",
        "Content-Type": "application/json",
    }

    http_status = None
    body: Dict[str, Any] = {}
    try:
        async with httpx.AsyncClient(timeout=25, follow_redirects=True) as client:
            resp = await client.post(
                "https://api.perplexity.ai/chat/completions",
                json=payload,
                headers=headers,
            )
            http_status = resp.status_code
            try:
                body = resp.json()
            except Exception:
                body = {}
            if resp.status_code != 200:
                err_obj = body.get("error") if isinstance(body.get("error"), dict) else None
                error_text = (err_obj or {}).get("message") or body.get("message") or f"Perplexity HTTP {resp.status_code}"
                if _scan_id:
                    await arecord_provider_trace(
                        scan_id=_scan_id, provider="perplexity", user_id=_user_id,
                        http_status=http_status,
                        latency_ms=int((_t.perf_counter() - _t0) * 1000),
                        request_summary=request_summary,
                        response_summary={"ok": False, "code": "PROVIDER_HTTP_ERROR"},
                        error=f"PROVIDER_HTTP_{http_status} : {str(error_text)[:200]}",
                        sanitiser_applied=True,
                    )
                return {"results": [], "error": str(error_text)[:240]}
    except Exception as exc:
        if _scan_id:
            await arecord_provider_trace(
                scan_id=_scan_id, provider="perplexity", user_id=_user_id,
                http_status=502,
                latency_ms=int((_t.perf_counter() - _t0) * 1000),
                request_summary=request_summary,
                response_summary={"ok": False, "code": "PROVIDER_NETWORK_ERROR"},
                error=f"PROVIDER_NETWORK_ERROR : {str(exc)[:200]}",
                sanitiser_applied=True,
            )
        return {"results": [], "error": str(exc)[:200]}

    content_text = ""
    try:
        content_text = (body.get("choices") or [{}])[0].get("message", {}).get("content", "") or ""
    except Exception:
        content_text = ""
    content_text = content_text.strip()
    if content_text.startswith("```"):
        content_text = re.sub(r"^```(?:json)?\s*", "", content_text)
        content_text = re.sub(r"\s*```$", "", content_text)

    citations = body.get("citations") if isinstance(body.get("citations"), list) else []

    parsed_results: List[Dict[str, Any]] = []
    try:
        import json as _json
        data = _json.loads(content_text) if content_text else []
        if isinstance(data, list):
            for i, item in enumerate(data[:num], start=1):
                if not isinstance(item, dict):
                    continue
                link = item.get("url") or item.get("link") or ""
                if not link:
                    continue
                parsed_results.append({
                    "title": str(item.get("title") or "")[:300],
                    "link": str(link)[:500],
                    "snippet": str(item.get("snippet") or item.get("description") or "")[:300],
                    "position": item.get("position") or i,
                })
    except Exception:
        # Structured-output parse failed — fall back to citations as bare links.
        # We still get useful provenance, just without titles/snippets.
        for i, url in enumerate(citations[:num], start=1):
            if not url:
                continue
            parsed_results.append({
                "title": "",
                "link": str(url)[:500],
                "snippet": "",
                "position": i,
            })

    if _scan_id:
        await arecord_provider_trace(
            scan_id=_scan_id, provider="perplexity", user_id=_user_id,
            http_status=200,
            latency_ms=int((_t.perf_counter() - _t0) * 1000),
            request_summary=request_summary,
            response_summary={"ok": True, "results_count": len(parsed_results)},
            evidence_payload={"results_count": len(parsed_results), "citations_count": len(citations)},
            sanitiser_applied=True,
        )
    return {"results": parsed_results, "error": None}


async def scrape_url_text(url: str) -> str:
    return await fetch_website_text(url)


# ==================== FILE PARSING ====================

def extract_text_from_pdf(file_content: bytes) -> str:
    if not PyPDF2:
        return "[PDF parsing not available]"
    try:
        pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
        text = ""
        for page in pdf_reader.pages[:20]:
            text += page.extract_text() or ""
        return text[:50000]
    except Exception as e:
        logger.error(f"PDF extraction error: {e}")
        return "[Could not extract PDF text]"


def extract_text_from_docx(file_content: bytes) -> str:
    if not DocxDocument:
        return "[DOCX parsing not available]"
    try:
        doc = DocxDocument(io.BytesIO(file_content))
        text = "\n".join([para.text for para in doc.paragraphs])
        return text[:50000]
    except Exception as e:
        logger.error(f"DOCX extraction error: {e}")
        return "[Could not extract DOCX text]"


def extract_text_from_xlsx(file_content: bytes) -> str:
    if not openpyxl:
        return "[XLSX parsing not available]"
    try:
        wb = openpyxl.load_workbook(io.BytesIO(file_content), data_only=True)
        text = ""
        for sheet in wb.worksheets[:5]:
            text += f"\n--- Sheet: {sheet.title} ---\n"
            for row in sheet.iter_rows(max_row=100, values_only=True):
                row_text = " | ".join([str(cell) if cell else "" for cell in row])
                if row_text.strip():
                    text += row_text + "\n"
        return text[:50000]
    except Exception as e:
        logger.error(f"XLSX extraction error: {e}")
        return "[Could not extract XLSX text]"


def extract_text_from_csv(file_content: bytes) -> str:
    try:
        text = file_content.decode('utf-8', errors='ignore')
        return text[:50000]
    except Exception as e:
        logger.error(f"CSV extraction error: {e}")
        return "[Could not extract CSV text]"


def extract_text_from_txt(file_content: bytes) -> str:
    try:
        return file_content.decode('utf-8', errors='ignore')[:50000]
    except Exception:
        return "[Could not extract text]"


async def extract_file_content(filename: str, file_content: bytes) -> str:
    """Extract text content from various file types."""
    ext = filename.lower().split('.')[-1] if '.' in filename else ''
    if ext == 'pdf':
        return extract_text_from_pdf(file_content)
    elif ext in ['docx', 'doc']:
        return extract_text_from_docx(file_content)
    elif ext in ['xlsx', 'xls']:
        return extract_text_from_xlsx(file_content)
    elif ext == 'csv':
        return extract_text_from_csv(file_content)
    elif ext in ['txt', 'md', 'json']:
        return extract_text_from_txt(file_content)
    else:
        return "[Unsupported file type for text extraction]"


# ==================== AUTH HELPERS ====================

def hash_password(password: str) -> str:
    return bcrypt.hashpw(password.encode('utf-8'), bcrypt.gensalt()).decode('utf-8')


def verify_password(password: str, hashed: str) -> bool:
    if hashed is None:
        return False
    return bcrypt.checkpw(password.encode('utf-8'), hashed.encode('utf-8'))


def create_token(user_id: str, email: str, role: str, account_id: Optional[str] = None) -> str:
    payload = {
        "sub": user_id,
        "email": email,
        "role": role,
        "account_id": account_id,
        "exp": datetime.now(timezone.utc) + timedelta(hours=JWT_EXPIRATION_HOURS)
    }
    return jwt.encode(payload, JWT_SECRET, algorithm=JWT_ALGORITHM)


def get_email_domain(email: str) -> str:
    if not email or "@" not in email:
        return ""
    return email.split("@", 1)[1].lower().strip()
